"""学习资料解析、清洗、分块与向量化入库（chromadb）。"""

import re
from pathlib import Path

import chromadb

from app.config import get_paths, get_settings
from app.kb.embeddings import embed_texts


def _read_txt(path: Path) -> str:
    """读取纯文本，utf-8 失败回退 gbk。"""
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="gbk")


def _read_pdf(path: Path) -> str:
    """逐页提取 PDF 文本。"""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(path: Path) -> str:
    """提取 docx 段落与表格单元格文本。"""
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _read_text(stored_path: str) -> str:
    """按扩展名提取文件纯文本。"""
    path = Path(stored_path)
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return _read_txt(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    raise ValueError(f"不支持的文件类型: {suffix or '(无扩展名)'}")


def _clean(text: str) -> str:
    """合并多余空白为单个空格。"""
    return re.sub(r"\s+", " ", text).strip()


def _chunk(text: str) -> list[str]:
    """字符滑动窗口分块，空块丢弃。"""
    size = get_settings().chunk_size
    step = size - get_settings().chunk_overlap
    chunks = []
    start = 0
    while start < len(text):
        piece = text[start : start + size].strip()
        if piece:
            chunks.append(piece)
        if start + size >= len(text):
            break
        start += step
    return chunks


def ingest_file(session_id: int, stored_path: str, filename: str, material_id: int) -> int:
    """解析→清洗→分块→向量化→写入 chroma，返回入库块数。"""
    chunks = _chunk(_clean(_read_text(stored_path)))
    if not chunks:
        return 0

    client = chromadb.PersistentClient(path=get_paths()["chroma_dir"])
    collection = client.get_or_create_collection(
        f"kb_{session_id}",
        metadata={"hnsw:space": "cosine"},
    )
    collection.add(
        ids=[f"s{session_id}_m{material_id}_i{i}" for i in range(len(chunks))],
        documents=chunks,
        metadatas=[
            {"session_id": session_id, "material_id": material_id, "filename": filename}
            for _ in chunks
        ],
        embeddings=embed_texts(chunks),
    )
    return len(chunks)
